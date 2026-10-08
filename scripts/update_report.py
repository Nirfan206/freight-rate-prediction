"""Append '17. Production Architecture and Future Deployment' to the existing DOCX report.

    python scripts/update_report.py

Idempotent (does nothing if the section exists). Never edits existing paragraphs/tables.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports" / "Spotter_Machine_Learning_Assessment_Report.docx"
HEADING = "17. Production Architecture and Future Deployment"

DISCLAIMER = ("The assessment solution is an offline batch ML pipeline. The additional production components "
              "provide a deployable architecture and local production simulation, but no public production "
              "deployment is claimed.")

COMPONENTS = [
    ("Prediction API (FastAPI)", "GET /health, GET /model/info, POST /predict. Accepts raw load features and runs validation, cleaning, feature engineering, the Ridge model and expm1 server-side, reusing the assessment cleaning and feature code. Errors are returned as structured JSON without stack traces."),
    ("Docker", "docker/Dockerfile (Python 3.13, non-root user, /health healthcheck) and docker-compose.yml for local development. No database, queue or cloud service is required."),
    ("Model registry", "File-based: models/registry/model_vN/{model.joblib, metadata.json} plus production.json. Versions are immutable. v1 is a byte-identical copy of the assessment artifact."),
    ("Monitoring", "Each prediction can be appended to a JSONL log (timestamp, load_id, prediction, interval, model version, latency). Summaries report counts, mean/median/min/max, distribution, latency and error count."),
    ("Drift detection", "Population Stability Index for numeric inputs and category-frequency PSI for equipment, against the training data. PSI < 0.10 stable, 0.10-0.25 warning, > 0.25 drift."),
    ("Performance monitoring", "When actual rates arrive: MAE, RMSE, MAPE and bias, grouped by equipment, distance bucket and month, compared with the holdout baseline. Without actuals the system reports 'actual target data unavailable'."),
    ("Retraining", "scripts/train_production.py runs validate, clean, features, temporal evaluation, alpha selection, final training and registration. It never promotes automatically; a policy check can recommend retraining from drift, degraded performance or model age."),
    ("Promotion and rollback", "scripts/promote_model.py and scripts/rollback_model.py validate the version (files, metadata, metrics, feature count, smoke prediction) and then rewrite production.json. Nothing is deleted."),
    ("Prediction intervals", "Empirical, residual-based intervals from the temporal-validation residuals on the log scale (multiplicative), stored with the model. They are an API-only addition; the official assessment predictions have no interval."),
    ("CI/CD", "GitHub Actions: tests.yml (imports, pytest, official scorer) and production-build.yml (tests, Docker build, container smoke check). No deployment step and no cloud credentials."),
]

FINDINGS = [
    "Prediction parity: the production path (raw input to prediction) reproduces the official validation_predictions.csv for all 12,000 rows to within 1.5e-10 when a model is trained with the same code.",
    "Alpha search: the production grid (0.001 to 100) reproduces the assessment result for alpha = 0.01 exactly (MAE 108.607337). The additional alpha = 0.001 is lower by 0.0006 MAE (about 0.0005%), a practical tie; the assessment model was not changed. A retraining run that selects strictly by MAE would therefore pick 0.001, and alpha can be pinned with --alphas 0.01.",
    "Interval coverage: the 90% nominal interval, calibrated on September 2025 residuals, covered 87.5% of October 2025 holdout rows. It is therefore slightly optimistic, which is why it is described as empirical and not calibrated.",
    "Drift: the assessment inputs (Nov-Dec 2025) show strong market_index drift against the training data (PSI above 3), consistent with the market-level shift already noted in the EDA. Other monitored features were stable.",
]

LIMITS = [
    "No public production deployment, load balancer, authentication or rate limiting is provided.",
    "Monitoring is batch and file based (JSONL plus scripts), not real-time streaming or alerting.",
    "No production prediction log or actual target values exist yet, so no production accuracy is claimed.",
    "Model files are pickles tied to the scikit-learn version that wrote them; only load files from a trusted registry.",
    "The API loads the model at startup; promotion or rollback requires an API restart.",
    "Prediction intervals are empirical, were slightly under-covering on the later holdout month, and are not calibrated probabilistic intervals.",
]


def main() -> int:
    doc = Document(str(REPORT))
    if any(p.text.strip() == HEADING for p in doc.paragraphs):
        print("Section already present; nothing to do.")
        return 0

    def sub(text):
        try:
            doc.add_heading(text, level=2)
        except KeyError:
            run = doc.add_paragraph().add_run(text)
            run.bold = True

    def bullets(items):
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()
    doc.add_heading(HEADING, level=1)
    para = doc.add_paragraph()
    para.add_run("Scope. ").bold = True
    para.add_run("Sections 1-16 describe the official assessment implementation and its results; they are unchanged. "
                 "This section describes additional components that are NOT required by the assessment.")
    para = doc.add_paragraph()
    para.add_run(DISCLAIMER).bold = True

    sub("17.1 Components")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Component", "What it does"
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for name, text in COMPONENTS:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, text
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)

    sub("17.2 Verified findings")
    bullets(FINDINGS)
    sub("17.3 Limitations")
    bullets(LIMITS)
    doc.save(str(REPORT))
    print(f"Appended section: {HEADING}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

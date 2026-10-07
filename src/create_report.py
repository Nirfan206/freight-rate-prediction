from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
REPORT_DIR.mkdir(exist_ok=True)

OUTPUT = REPORT_DIR / "Spotter_Machine_Learning_Assessment_Report.docx"
CHART = ROOT / "scorer_results" / "candidate_december.png"


def set_cell_shading(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


doc = Document()

# Margins
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Default font
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run("Spotter Machine Learning Assessment")
run.bold = True
run.font.size = Pt(22)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run("Freight Rate Prediction")
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Machine Learning Engineer Take-Home Assessment").italic = True

doc.add_paragraph("")

# Executive Summary
add_heading(doc, "1. Executive Summary")

doc.add_paragraph(
    "This project develops a machine learning system for predicting posted freight "
    "rates from historical load information. The development dataset contains 48,000 "
    "historical loads from January through October 2025, while the assessment "
    "validation dataset contains 12,000 loads from November through December 2025."
)

doc.add_paragraph(
    "The workflow covers data auditing, cleaning, exploratory data analysis, feature "
    "engineering, temporal validation, model comparison, hyperparameter tuning, "
    "error analysis, final model training, validation prediction generation, and "
    "December scenario prediction."
)

doc.add_paragraph(
    "The final selected model is Ridge Regression trained on the log-transformed "
    "target using alpha = 0.01. On the temporal validation split covering September "
    "and October 2025, it achieved an MAE of approximately $108.61, RMSE of "
    "$633.26, and R² of 0.8278."
)

# Dataset
add_heading(doc, "2. Dataset")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"

headers = ["Dataset", "Rows", "Purpose"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    set_cell_shading(table.rows[0].cells[i])

rows = [
    ("train-test.csv", "48,000", "Labeled development/training data"),
    ("validation.csv", "12,000", "Final assessment prediction data"),
    ("validation-predictions-template.csv", "12,000", "Required prediction format"),
    ("december-chart-inputs.csv", "31", "Fixed December scenario"),
]

for dataset, count, purpose in rows:
    cells = table.add_row().cells
    cells[0].text = dataset
    cells[1].text = count
    cells[2].text = purpose

doc.add_paragraph(
    "The training data contains 14 columns including the target posted_rate. "
    "The assessment validation data contains 13 input columns because the target "
    "is not provided."
)

# Data audit
add_heading(doc, "3. Data Audit")

add_bullet(doc, "Training rows: 48,000")
add_bullet(doc, "Training columns: 14")
add_bullet(doc, "Validation rows: 12,000")
add_bullet(doc, "Exact duplicate training rows: 0")
add_bullet(doc, "Exact duplicate validation rows: 0")
add_bullet(doc, "Duplicate training load IDs: 0")
add_bullet(doc, "Duplicate validation load IDs: 0")
add_bullet(doc, "Training missing weight values: 300")
add_bullet(doc, "Training negative weight values: 292")
add_bullet(doc, "Training missing market_index values: 374")
add_bullet(doc, "Validation missing weight values: 165")
add_bullet(doc, "Validation negative weight values: 145")
add_bullet(doc, "Validation missing market_index values: 249")

doc.add_paragraph(
    "The target posted_rate ranges from approximately $57 to $25,533. "
    "The target distribution is right-skewed, motivating the use of a log-transformed "
    "target for the Ridge model."
)

# Cleaning
add_heading(doc, "4. Data Cleaning")

doc.add_paragraph(
    "Cleaning statistics were fitted using training data only to reduce the risk "
    "of validation leakage."
)

add_number(
    doc,
    "Missing and invalid weight values were identified using explicit diagnostic flags."
)
add_number(
    doc,
    "Missing weight values were imputed using the training-set median."
)
add_number(
    doc,
    "Weight-related derived features were created while preserving raw-weight information."
)
add_number(
    doc,
    "Missing market_index values were first handled using date-specific training "
    "statistics and then a training global fallback."
)
add_number(
    doc,
    "Distance values were checked for validity before feature generation."
)

# EDA
add_heading(doc, "5. Exploratory Data Analysis")

doc.add_paragraph(
    "Exploratory analysis examined the target distribution, rate per mile, distance, "
    "weight, equipment type, market index, quote signal, temporal patterns, and "
    "relationships between the variables."
)

add_bullet(
    doc,
    "Distance has a very strong relationship with posted_rate; the Spearman correlation "
    "was approximately 0.976."
)
add_bullet(
    doc,
    "Rate per mile generally decreases as shipment distance increases."
)
add_bullet(
    doc,
    "Median rate per mile differs by equipment type, with Reefer and Flatbed generally "
    "above Dry Van."
)
add_bullet(
    doc,
    "Market index and quote signal provide additional information beyond distance."
)
add_bullet(
    doc,
    "A small number of extreme target values create substantial RMSE impact."
)

# Feature engineering
add_heading(doc, "6. Feature Engineering")

doc.add_paragraph(
    "The final modeling pipeline generates 43 model features."
)

features = [
    "Date-derived features such as day of week, day of month, and day of year",
    "Log-transformed distance and weight-related features",
    "Weight-distance interaction features",
    "Latitude and longitude differences",
    "Absolute geographic differences",
    "Haversine distance",
    "Distance-to-geographic-distance relationships",
    "Market-index and distance interactions",
    "Equipment indicator variables",
    "Equipment-specific distance and market interactions",
]

for feature in features:
    add_bullet(doc, feature)

doc.add_paragraph(
    "Identifiers such as load_id were excluded from model features. Target-derived "
    "variables such as posted_rate and rate_per_mile were also excluded."
)

# Validation
add_heading(doc, "7. Validation Strategy")

doc.add_paragraph(
    "Because freight rates can change over time, a temporal validation strategy was "
    "used instead of a random split."
)

add_bullet(doc, "Training period: January 1 through August 31, 2025")
add_bullet(doc, "Validation period: September 1 through October 31, 2025")
add_bullet(doc, "Temporal training rows: 38,477")
add_bullet(doc, "Temporal validation rows: 9,523")

doc.add_paragraph(
    "This design evaluates whether the model can generalize from earlier market "
    "conditions to later unseen dates."
)

# Experiments
add_heading(doc, "8. Model Experiments")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"

for i, h in enumerate(["Model", "MAE", "RMSE", "R²"]):
    table.rows[0].cells[i].text = h
    set_cell_shading(table.rows[0].cells[i])

results = [
    ("Ridge Log Target", "108.61", "633.26", "0.8278"),
    ("HistGradientBoosting", "135.44", "635.92", "0.8264"),
    ("Extra Trees", "141.79", "644.40", "0.8217"),
    ("Random Forest", "158.65", "648.79", "0.8193"),
    ("Equipment Rate/Mile", "229.10", "670.66", "0.8069"),
    ("Global Rate/Mile", "256.95", "684.25", "0.7990"),
    ("Global Median", "1148.92", "1569.42", "-0.0577"),
]

for model, mae, rmse, r2 in results:
    cells = table.add_row().cells
    cells[0].text = model
    cells[1].text = mae
    cells[2].text = rmse
    cells[3].text = r2

doc.add_paragraph(
    "Ridge Regression with a log-transformed target achieved the lowest MAE among "
    "the evaluated approaches and was selected for further tuning."
)

# Tuning
add_heading(doc, "9. Hyperparameter Tuning")

doc.add_paragraph(
    "The Ridge regularization parameter alpha was evaluated across multiple values."
)

add_bullet(doc, "Best alpha: 0.01")
add_bullet(doc, "Validation MAE: $108.61")
add_bullet(doc, "Validation RMSE: $633.26")
add_bullet(doc, "Validation R²: 0.8278")

doc.add_paragraph(
    "The difference between alpha = 0.01 and alpha = 0.1 was very small, while "
    "larger regularization values progressively degraded validation performance."
)

# Error analysis
add_heading(doc, "10. Error Analysis")

doc.add_paragraph(
    "The tuned Ridge model achieved an overall temporal-validation MAPE of "
    "approximately 4.81%."
)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"

for i, h in enumerate(["Equipment", "Rows", "MAE"]):
    table.rows[0].cells[i].text = h
    set_cell_shading(table.rows[0].cells[i])

for equipment, rows_count, mae in [
    ("Dry Van", "5,360", "$103.54"),
    ("Flatbed", "1,770", "$111.01"),
    ("Reefer", "2,393", "$118.19"),
]:
    cells = table.add_row().cells
    cells[0].text = equipment
    cells[1].text = rows_count
    cells[2].text = mae

doc.add_paragraph("Distance-based error analysis showed:")
add_bullet(doc, "≤200 miles: MAE approximately $18.99")
add_bullet(doc, "201–500 miles: MAE approximately $51.57")
add_bullet(doc, "501–1,000 miles: MAE approximately $72.28")
add_bullet(doc, "1,001–1,500 miles: MAE approximately $124.35")
add_bullet(doc, "1,501–2,000 miles: MAE approximately $180.24")
add_bullet(doc, ">2,000 miles: MAE approximately $191.38")

doc.add_paragraph(
    "The largest errors were associated with extreme posted-rate observations. "
    "These observations strongly influence RMSE even when the majority of predictions "
    "have substantially smaller absolute errors."
)

# Final model
add_heading(doc, "11. Final Model")

doc.add_paragraph(
    "After model selection and tuning, the final Ridge model with alpha = 0.01 "
    "was retrained on all 48,000 development rows."
)

add_bullet(doc, "Model: Ridge Regression")
add_bullet(doc, "Target transformation: log1p(posted_rate)")
add_bullet(doc, "Alpha: 0.01")
add_bullet(doc, "Model features: 43")
add_bullet(doc, "Training rows: 48,000")
add_bullet(doc, "Artifact: artifacts/ridge_model.joblib")

# Predictions
add_heading(doc, "12. Assessment Predictions")

doc.add_paragraph(
    "The final model generated predictions for all 12,000 rows in validation.csv."
)

add_bullet(doc, "Output file: validation_predictions.csv")
add_bullet(doc, "Rows: 12,000")
add_bullet(doc, "Required columns: load_id, predicted_rate")
add_bullet(doc, "IDs match TE-000001 through TE-012000")
add_bullet(doc, "All predicted rates are finite and positive")

# December
add_heading(doc, "13. Fixed December Scenario")

doc.add_paragraph(
    "Predictions were also generated for all 31 December scenario rows. "
    "The supplied December input does not contain market_index or quote_signal, "
    "so the current implementation uses training-set median fallback values for "
    "these fields."
)

add_bullet(doc, "December rows: 31")
add_bullet(doc, "Scenario: Lexington to Fort Wayne")
add_bullet(doc, "Distance: 360 miles")
add_bullet(doc, "Equipment: Dry Van")
add_bullet(doc, "Weight: 32,000")
add_bullet(doc, "Market index fallback: training median")
add_bullet(doc, "Quote signal fallback: training median")

if CHART.exists():
    doc.add_paragraph("December prediction chart:")
    doc.add_picture(str(CHART), width=Inches(6.3))

# Scorer
add_heading(doc, "14. Official Scorer Verification")

doc.add_paragraph(
    "The provided assessment scorer was executed using the generated validation "
    "and December predictions."
)

p = doc.add_paragraph()
run = p.add_run("Result: PASSED")
run.bold = True

add_bullet(doc, "Validated 12,000 final predictions.")
add_bullet(doc, "Validated 31 fixed December predictions.")
add_bullet(doc, "Created scorer_results/candidate_december.png.")

doc.add_paragraph(
    "The hidden final Spotter evaluation metric is not calculated by the provided "
    "local scorer and will be determined by Spotter after submission."
)

# Limitations
add_heading(doc, "15. Limitations and Future Improvements")

add_bullet(
    doc,
    "The temporal validation split does not contain the new-city cases present in "
    "the final assessment validation set."
)
add_bullet(
    doc,
    "Extreme freight-rate outliers remain difficult for the current linear model."
)
add_bullet(
    doc,
    "The December input lacks some features available in the historical development data."
)
add_bullet(
    doc,
    "Additional location-based features, robust regression, gradient boosting, "
    "and carefully designed target-encoding strategies could be investigated."
)
add_bullet(
    doc,
    "Rolling-origin validation across multiple time windows could provide a more "
    "robust estimate of temporal generalization."
)

# Conclusion
add_heading(doc, "16. Conclusion")

doc.add_paragraph(
    "The completed pipeline provides a reproducible end-to-end solution for freight "
    "rate prediction. The workflow separates training-time preprocessing from "
    "inference, uses a time-aware validation strategy, compares multiple modeling "
    "approaches, tunes the selected model, analyzes prediction errors, generates the "
    "required assessment predictions, and successfully passes the provided scorer."
)

doc.add_paragraph(
    "The selected Ridge model provides the strongest validation MAE among the tested "
    "approaches and was retrained on the complete development dataset before producing "
    "the final assessment predictions."
)

doc.save(OUTPUT)

print(f"Created report: {OUTPUT}")